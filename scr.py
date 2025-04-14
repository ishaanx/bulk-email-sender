import sys
import smtplib
import time
import os
from datetime import datetime
from openpyxl import load_workbook
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLineEdit, QTextEdit, QPushButton, QProgressBar,
    QVBoxLayout, QHBoxLayout, QFileDialog, QMessageBox, QFormLayout, QComboBox
)
from PySide6.QtGui import QAction, QIcon
from PySide6.QtCore import QThread, Signal, Qt

# Load recipients from Excel
def load_recipients(file_path):
    wb = load_workbook(filename=file_path)
    ws = wb.active
    return [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]


# Send email
def send_email(smtp_server, smtp_port, sender_email, sender_password, to_email, subject, body, format_type):
    msg = MIMEMultipart()
    msg['From'] = sender_email
    msg['To'] = to_email
    msg['Subject'] = subject
    msg.attach(MIMEText(body, format_type))

    with smtplib.SMTP(smtp_server, int(smtp_port)) as server:
        server.starttls()
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, to_email, msg.as_string())


# Threaded sender
class EmailSenderThread(QThread):
    progress = Signal(int)
    finished = Signal(int, int)
    error = Signal(str)

    def __init__(self, smtp_server, smtp_port, sender_email, password, subject, body, recipients, log_file_path, format_type):
        super().__init__()
        self.smtp_server = smtp_server
        self.smtp_port = smtp_port
        self.sender_email = sender_email
        self.password = password
        self.subject = subject
        self.body = body
        self.recipients = recipients
        self.log_file_path = log_file_path
        self.format_type = format_type
        self.success_count = 0
        self.failure_count = 0

    def run(self):
        try:
            with open(self.log_file_path, 'a', encoding='utf-8') as log_file:
                for i, email in enumerate(self.recipients):
                    timestamp = self.timestamp()
                    try:
                        log_file.write(f"{timestamp} Sending to {email}...\n")
                        send_email(
                            self.smtp_server, self.smtp_port,
                            self.sender_email, self.password,
                            email, self.subject, self.body,
                            self.format_type
                        )
                        log_file.write(f"{timestamp} ✓ Sent to {email}\n")
                        self.success_count += 1
                    except Exception as e:
                        log_file.write(f"{timestamp} ✗ Failed to send to {email}: {str(e)}\n")
                        self.failure_count += 1

                    self.progress.emit(int(((i + 1) / len(self.recipients)) * 100))
                    if i < len(self.recipients) - 1:
                        time.sleep(30)

            self.finished.emit(self.success_count, self.failure_count)

        except Exception as e:
            self.error.emit(str(e))

    def timestamp(self):
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# GUI app
class EmailSenderApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("📧 Bulk Email Sender (PySide6)")
        self.setFixedSize(500, 450)
        self.setWindowIcon(QIcon("email_icon.ico"))

        self.latest_log_file = None
        self.is_dark_mode = False

        # Central widget
        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        # Input fields
        self.smtp_input = QLineEdit("smtp.gmail.com")
        self.port_input = QLineEdit("587")
        self.email_input = QLineEdit()
        self.password_input = QLineEdit()
        self.password_input.setEchoMode(QLineEdit.Password)

        self.file_input = QLineEdit()
        browse_button = QPushButton("Browse")
        browse_button.clicked.connect(self.browse_file)

        file_layout = QHBoxLayout()
        file_layout.addWidget(self.file_input)
        file_layout.addWidget(browse_button)

        self.subject_input = QLineEdit()
        self.body_input = QTextEdit()

        self.format_combo = QComboBox()
        self.format_combo.addItems(["Plain Text", "HTML"])

        self.progress = QProgressBar()
        self.send_button = QPushButton("Send Emails")
        self.send_button.clicked.connect(self.send_emails)

        # Form layout
        form = QFormLayout()
        form.addRow("SMTP Server:", self.smtp_input)
        form.addRow("Port:", self.port_input)
        form.addRow("Your Email:", self.email_input)
        form.addRow("Password:", self.password_input)
        form.addRow("Excel File:", file_layout)
        form.addRow("Subject:", self.subject_input)
        form.addRow("Body:", self.body_input)
        form.addRow("Format:", self.format_combo)

        # Main layout
        layout = QVBoxLayout()
        layout.addLayout(form)
        layout.addWidget(self.progress)
        layout.addWidget(self.send_button)
        central_widget.setLayout(layout)

        # Menu
        self.setup_menu()

    def setup_menu(self):
        menubar = self.menuBar()

        logs_menu = menubar.addMenu("Logs")
        open_log_action = QAction("Open Last Log File", self)
        open_log_action.triggered.connect(self.open_last_log_file)
        logs_menu.addAction(open_log_action)

        view_menu = menubar.addMenu("View")
        toggle_theme_action = QAction("Toggle Dark Mode", self)
        toggle_theme_action.triggered.connect(self.toggle_dark_mode)
        view_menu.addAction(toggle_theme_action)

        help_menu = menubar.addMenu("Help")
        about_action = QAction("About", self)
        about_action.triggered.connect(self.show_about_dialog)
        help_menu.addAction(about_action)

        exit_action = QAction("Exit", self)
        exit_action.triggered.connect(self.close)
        menubar.addAction(exit_action)

    def browse_file(self):
        file_path, _ = QFileDialog.getOpenFileName(self, "Select Excel File", "", "Excel Files (*.xlsx)")
        if file_path:
            self.file_input.setText(file_path)

    def send_emails(self):
        file_path = self.file_input.text().strip()
        try:
            recipients = load_recipients(file_path)
            if not recipients:
                raise ValueError("No recipients found.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Failed to load recipients: {e}")
            return

        self.send_button.setEnabled(False)
        self.progress.setValue(0)

        self.latest_log_file = f"email_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        format_type = "html" if self.format_combo.currentText() == "HTML" else "plain"

        self.thread = EmailSenderThread(
            smtp_server=self.smtp_input.text(),
            smtp_port=self.port_input.text(),
            sender_email=self.email_input.text(),
            password=self.password_input.text(),
            subject=self.subject_input.text(),
            body=self.body_input.toPlainText(),
            recipients=recipients,
            log_file_path=self.latest_log_file,
            format_type=format_type
        )
        self.thread.progress.connect(self.progress.setValue)
        self.thread.finished.connect(self.on_finished)
        self.thread.error.connect(self.on_error)
        self.thread.start()

    def on_finished(self, success, failure):
        QMessageBox.information(self, "Done",
            f"Emails Sent: {success}\nFailed: {failure}\n\nLog file saved as:\n{self.latest_log_file}")
        self.send_button.setEnabled(True)

    def on_error(self, message):
        QMessageBox.critical(self, "Error", message)
        self.send_button.setEnabled(True)

    def open_last_log_file(self):
        if self.latest_log_file and os.path.exists(self.latest_log_file):
            os.startfile(self.latest_log_file)
        else:
            QMessageBox.warning(self, "No Log File", "No log file available to open.")

    def show_about_dialog(self):
        QMessageBox.information(
            self,
            "About",
            "📧 Bulk Email Sender with:\n• HTML/plaintext format\n• Logging\n• Dark Mode\n• v0.0.9\n\nCreated by ishaanx"
        )

    def toggle_dark_mode(self):
        if not self.is_dark_mode:
            dark_stylesheet = """
                QWidget { background-color: #2b2b2b; color: #f0f0f0; }
                QLineEdit, QTextEdit, QComboBox, QProgressBar { background-color: #3c3c3c; border: 1px solid #555; }
                QPushButton { background-color: #444; color: #fff; }
                QPushButton:hover { background-color: #666; }
            """
            self.setStyleSheet(dark_stylesheet)
        else:
            self.setStyleSheet("")
        self.is_dark_mode = not self.is_dark_mode


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = EmailSenderApp()
    window.show()
    sys.exit(app.exec())
